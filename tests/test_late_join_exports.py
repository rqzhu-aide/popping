"""Late team joining preserves historical activity and the export layout."""

from datetime import datetime
import io
import zipfile

from openpyxl import load_workbook

from test_workflow_safety import (
    _activate_current_schema_presentation,
    _instructor_client,
    _participant_rows,
    _post_after_rating_grace,
    _set_state,
    _state_row,
    _student_client,
    app_module,
    course_env,
)


def _export_workbook(instructor, env):
    response = instructor.get(f"/export/{env['slug']}")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        return load_workbook(io.BytesIO(archive.read("course_data.xlsx")))


def _rows(workbook, sheet_name):
    values = list(workbook[sheet_name].iter_rows(values_only=True))
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_late_join_exports_only_actual_activity_without_changing_layout(
        course_env, monkeypatch):
    env = course_env
    clock = {"now": datetime(2026, 9, 8, 12, 0, 0)}
    monkeypatch.setattr(app_module, "_utcnow", lambda: clock["now"])
    instructor = _instructor_client(env)
    baseline = _export_workbook(instructor, env)
    headers = {
        name: next(baseline[name].iter_rows(values_only=True))
        for name in baseline.sheetnames
    }
    baseline.close()

    _set_state(env, phase="discussion")
    earlier_thumb = _student_client(env, "s1").post(
        "/api/grade_peer", json={"recipient_id": "s2", "selected": True},
    )
    assert earlier_thumb.status_code == 200

    def finish_presentation(key, team_name, rater=None):
        history = _state_row(env)["presentation_history"]
        timestamp = clock["now"].strftime("%Y-%m-%d %H:%M:%S.%f")
        _activate_current_schema_presentation(
            env,
            active_team_id=env["teams"][team_name],
            poll_question_key=key,
            presentation_started_at=timestamp,
            presentation_created_at=timestamp,
            presentation_history=history,
            poll_active=1 if rater else 0,
            poll_started_at=timestamp if rater else None,
        )
        if rater:
            rating = _student_client(env, rater).post(
                "/api/submit_rating",
                json={
                    "presentation_key": key,
                    "q1_developed": 4,
                    "q2_easy": 5,
                },
            )
            assert rating.status_code == 200
            stopped = instructor.post(
                "/api/stop_poll", json={"presentation_key": key},
            )
            assert stopped.status_code == 200
            finished = _post_after_rating_grace(
                instructor, "/api/next_presentation",
                {"presentation_key": key}, clock,
            )
        else:
            finished = instructor.post(
                "/api/next_presentation", json={"presentation_key": key},
            )
        assert finished.status_code == 200

    finish_presentation("pres-before-join", "Team 1", rater="s4")
    earlier_participants = _participant_rows(env)
    assert {row["student_identifier"] for row in earlier_participants} == {
        "s1", "s2",
    }

    late_student = _student_client(env, "s3")
    joined = late_student.post(
        "/api/join_team", json={"team_id": env["teams"]["Team 1"]},
    )
    assert joined.status_code == 200
    assert _participant_rows(env) == earlier_participants

    _set_state(env, phase="discussion")
    later_thumb = late_student.post(
        "/api/grade_peer", json={"recipient_id": "s1", "selected": True},
    )
    assert later_thumb.status_code == 200
    finish_presentation("pres-after-join-rated", "Team 2", rater="s3")
    finish_presentation("pres-after-join-presented", "Team 1")
    _set_state(env, phase="ended")

    workbook = _export_workbook(instructor, env)
    try:
        assert workbook.sheetnames == list(headers)
        for name, header in headers.items():
            assert next(workbook[name].iter_rows(values_only=True)) == header

        participants = _rows(workbook, "Presentation Participants")
        assert {
            row["participant_id"] for row in participants
            if row["presentation_key"] == "pres-before-join"
        } == {"s1", "s2"}
        assert [
            row["presentation_key"] for row in participants
            if row["participant_id"] == "s3"
        ] == ["pres-after-join-presented"]

        ratings = _rows(workbook, "Presentation Ratings")
        assert {
            (row["presentation_key"], row["rater_id"], row["rater_team"])
            for row in ratings
        } == {
            ("pres-before-join", "s4", "Team 2"),
            ("pres-after-join-rated", "s3", "Team 1"),
        }
        assert {
            (row["grader_id"], row["recipient_id"], row["grader_team"])
            for row in _rows(workbook, "Peer Reviews")
        } == {("s1", "s2", "Team 1"), ("s3", "s1", "Team 1")}

        students = _rows(workbook, "Students")
        assert {row["student_id"] for row in students} == {
            "s1", "s2", "s3", "s4",
        }
        late_row = next(row for row in students if row["student_id"] == "s3")
        assert late_row["team"] == "Team 1"
        assert late_row["course_presentation_team_turns"] == 1
        assert late_row["course_challenger_turns"] == 0
        assert late_row["thumbs_given"] == 1
        assert late_row["thumbs_received"] == 0
        assert late_row["presentation_ratings_given"] == 1
        assert late_row["challenges_rated"] == 0
    finally:
        workbook.close()
